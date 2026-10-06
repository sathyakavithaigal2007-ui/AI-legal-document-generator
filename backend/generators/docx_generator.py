import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from typing import Dict, Any

def generate_docx_bytes(doc_data: Dict[str, Any]) -> bytes:
    """Generates a professionally styled Microsoft Word (.docx) file."""
    doc = Document()
    
    # Configure 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title_text = doc_data.get("title", "LEGAL AGREEMENT")
    content = doc_data.get("content", "")
    branding = doc_data.get("branding", {}) or {}
    company_name = branding.get("company_name", "")
    footer_text = branding.get("footer_text", "Generated via LegalEase AI Document Generator")
    author = branding.get("author_name", "")

    # Header section if company name provided
    if company_name:
        header_para = doc.add_paragraph()
        header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_comp = header_para.add_run(company_name.upper())
        run_comp.font.name = 'Arial'
        run_comp.font.size = Pt(9)
        run_comp.font.bold = True
        run_comp.font.color.rgb = RGBColor(99, 102, 241) # Indigo accent

    # Document Main Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Pt(12)
    title_para.paragraph_format.space_after = Pt(18)
    
    title_run = title_para.add_run(title_text.upper())
    title_run.font.name = 'Times New Roman'
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Deep navy

    # Add subtitle line if effective date
    if doc_data.get("effective_date"):
        sub_para = doc.add_paragraph()
        sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_para.paragraph_format.space_after = Pt(24)
        sub_run = sub_para.add_run(f"EFFECTIVE DATE: {doc_data.get('effective_date')}")
        sub_run.font.name = 'Times New Roman'
        sub_run.font.size = Pt(10)
        sub_run.font.italic = True

    # Parse and add content lines
    lines = content.split('\n')
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15

        # Check if heading (e.g. "1. DEFINITIONS", "PARTIES:", "IN WITNESS WHEREOF")
        if stripped.isupper() or stripped.endswith(":") or (len(stripped) > 2 and stripped[0].isdigit() and stripped[1] in ['.', ' ']):
            run = p.add_run(stripped)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
            if stripped.isupper() and len(stripped) < 40:
                p.paragraph_format.space_before = Pt(12)
        elif stripped.startswith("BY:") or stripped.startswith("TITLE:") or stripped.startswith("ENTITY:") or stripped.startswith("DATE:"):
            run = p.add_run(stripped)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = True
        elif stripped.startswith("___"):
            run = p.add_run(stripped)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            p.paragraph_format.space_before = Pt(16)
        else:
            run = p.add_run(stripped)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 41, 59)

    # Footer
    section = doc.sections[0]
    footer = section.footer
    f_para = footer.paragraphs[0]
    f_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = f_para.add_run(f"{footer_text}  •  Page 1")
    f_run.font.name = 'Arial'
    f_run.font.size = Pt(8.5)
    f_run.font.color.rgb = RGBColor(100, 116, 139)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
